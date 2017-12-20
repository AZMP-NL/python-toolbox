'''
From https://pypi.python.org/pypi/pyshp:

A record in a shapefile contains the attributes for each shape in the
collection of geometry. Records are stored in the dbf file. The link between
geometry and attributes is the foundation of all geographic information systems.
This critical link is implied by the order of shapes and corresponding records
in the shp geometry file and the dbf attribute file.

The field names of a shapefile are available as soon as you read a shapefile.
You can call the "fields" attribute of the shapefile as a Python list. Each
field is a Python list with the following information:

* Field name: the name describing the data at this column index.
* Field type: the type of data at this column index. Types can be: Character,
Numbers, Longs, Dates, or Memo. The "Memo" type has no meaning within a
GIS and is part of the xbase spec instead.
* Field length: the length of the data found at this column index. Older GIS
software may truncate this length to 8 or 11 characters for "Character"
fields.
* Decimal length: the number of decimal places found in "Number" fields.

To see the fields for the Reader object above (sf) call the "fields"
attribute:
'''


## from netCDF4 import Dataset, num2date
## import time, calendar, datetime, numpy
## from mpl_toolkits.basemap import Basemap
## import matplotlib.pyplot as plt
## import urllib, os
## # data downloaded from the form at
## # http://coastwatch.pfeg.noaa.gov/erddap/tabledap/apdrcArgoAll.html
## filename, headers = urllib.urlretrieve('http://coastwatch.pfeg.noaa.gov/erddap/tabledap/apdrcArgoAll.nc?longitude,latitude,time&longitude>=0&longitude<=360&latitude>=-90&latitude<=90&time>=2010-01-01&time<=2010-01-08&distinct()')
## dset = Dataset(filename)
## lats = dset.variables['latitude'][:]
## lons = dset.variables['longitude'][:]
## time = dset.variables['time']
## times = time[:]
## t1 = times.min(); t2 = times.max()
## date1 = num2date(t1, units=time.units)
## date2 = num2date(t2, units=time.units)
## dset.close()
## os.remove(filename)
## # draw map with markers for float locations
## m = Basemap(projection='hammer',lon_0=180)
## x, y = m(lons,lats)
## m.drawmapboundary(fill_color='#99ffff')
## m.fillcontinents(color='#cc9966',lake_color='#99ffff')
## m.scatter(x,y,3,marker='o',color='k')
## plt.title('Locations of %s ARGO floats active between %s and %s' %\
##         (len(lats),date1,date2),fontsize=12)
## plt.show()


import netCDF4
from mpl_toolkits.basemap import Basemap
import numpy as  np
import matplotlib.pyplot as plt
import openpyxl, pprint
import shapefile 

## ---- Region parameters ---- ##
dataFile = '/home/cyrf0006/Data/GEBCO/GRIDONE_1D.nc'
lon_0 = -50
lat_0 = 50
lonLims = [-70, -40]
latLims = [40, 65]
proj = 'merc'
decim_scale = 4
stationFile = '/home/cyrf0006/research/AZMP_surveys/STANDARD_SECTIONS.xlsx'
fig_name = 'AZMP_NAFO.png'
ephem = 'ephem_calval.txt'
swot_kml = 'SWOT_Science_sept2015_Swath_10_60.kml'

## ---- Bathymetry ---- ####
v = np.linspace(-4000, 0, 9)

# Load data
dataset = netCDF4.Dataset(dataFile)

# Extract variables
x = dataset.variables['x_range']
y = dataset.variables['y_range']
spacing = dataset.variables['spacing']

# Compute Lat/Lon
nx = int((x[-1]-x[0])/spacing[0]) + 1  # num pts in x-dir
ny = int((y[-1]-y[0])/spacing[1]) + 1  # num pts in y-dir

lon = np.linspace(x[0],x[-1],nx)
lat = np.linspace(y[0],y[-1],ny)

# Reshape data
zz = dataset.variables['z']
Z = zz[:].reshape(ny, nx)

# Reduce data according to Region params
lon = lon[::decim_scale]
lat = lat[::decim_scale]
Z = Z[::decim_scale, ::decim_scale]

## ---- Station info ---- ##
import pandas as pd
df = pd.read_excel(stationFile)
#print the column names
print df.columns
#get the values for a given column
sections = df['SECTION'].values
stations = df['STATION'].values
stationLat = df['LAT'].values
stationLon = df['LONG.1'].values

index_SEGB = df.SECTION[df.SECTION=="SOUTHEAST GRAND BANK"].index.tolist()
index_FC = df.SECTION[df.SECTION=="FLEMISH CAP"].index.tolist()
index_BB = df.SECTION[df.SECTION=="BONAVISTA"].index.tolist()
index_WB = df.SECTION[df.SECTION=="WHITE BAY"].index.tolist()
index_SI = df.SECTION[df.SECTION=="SEAL ISLAND"].index.tolist()
index_MB = df.SECTION[df.SECTION=="MAKKOVIK BANK"].index.tolist()
index_BI = df.SECTION[df.SECTION=="BEACH ISLAND"].index.tolist()
index_FI = df.SECTION[df.SECTION=="FUNK ISLAND"].index.tolist()
index_S27 = df.SECTION[df.SECTION=="STATION 27"].index.tolist()
index_SESPB = df.SECTION[df.SECTION=="SOUTHEAST ST PIERRE BANK"].index.tolist()
index_SWSPB = df.SECTION[df.SECTION=="SOUTHWEST ST PIERRE BANK"].index.tolist()
index_SS = df.SECTION[df.SECTION=="SMITH SOUND"].index.tolist()

## ---- NAFO divisions ---- ##
## shapef = '/home/cyrf0006/research/AZMP_utils/NAFO_divisions/Divisions/Divisions.shp'
## sf = shapefile.Reader(shapef)
## shapes = sf.shapes()
## records = sf.records()
## sf.fields

## shapes = sf.shapes()
## len(list(sf.iterShapes()))

# OR
myshp = open('/home/cyrf0006/research/AZMP_utils/NAFO_divisions/Divisions/Divisions.shp', 'rb')
mydbf = open('/home/cyrf0006/research/AZMP_utils/NAFO_divisions/Divisions/Divisions.dbf', 'rb')
r = shapefile.Reader(shp=myshp, dbf=mydbf)
records = r.records()
shapes = r.shapes()



## ---- plot ---- ##
fig = plt.figure(1)
#m = Basemap(projection='ortho',lon_0=lon_0,lat_0=lat_0,resolution=None)
m = Basemap(projection='merc',lon_0=lon_0,lat_0=lat_0, llcrnrlon=lonLims[0],llcrnrlat=latLims[0],urcrnrlon=lonLims[1],urcrnrlat=latLims[1], resolution='l')
x,y = m(*np.meshgrid(lon,lat))
c = m.contourf(x, y, np.flipud(Z), v, cmap=plt.cm.PuBu_r, extend="min");
#c = m.contourf(x, y, np.flipud(Z), v, cmap=plt.cm.PuRd_r, extend="min");
m.fillcontinents(color='grey');
m.drawparallels(np.arange(10,70,10), labels=[1,0,0,0], fontsize=12, fontweight='bold');
m.drawmeridians(np.arange(-80, 5, 10), labels=[0,0,0,1], fontsize=12, fontweight='bold');
cb = plt.colorbar(c)
for l in cb.ax.yaxis.get_ticklabels():
    l.set_weight("bold")
    l.set_fontsize(10)
cb.set_label('Depth (m)', fontsize=13, fontweight='bold')
plt.title("AZMP-NL standard lines", fontsize=13, fontweight='bold')

# plot stations
x, y = m(stationLon[index_SEGB],stationLat[index_SEGB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' SEGB', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_FC],stationLat[index_FC])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' FC', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_BB],stationLat[index_BB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' BB', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_WB],stationLat[index_WB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' WB', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_SI],stationLat[index_SI])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' SI', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_MB],stationLat[index_MB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' MB', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_BI],stationLat[index_BI])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], ' BI', horizontalalignment='left', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_FI],stationLat[index_FI])
m.scatter(x,y,3,marker='o',color='lightcoral')
plt.text(x[-1], y[-1], ' FI', horizontalalignment='left', verticalalignment='center', fontsize=10, color='lightcoral', fontweight='bold')
x, y = m(stationLon[index_S27],stationLat[index_S27])
m.scatter(x,y,3,marker='o',color='lightcoral')
#m.scatter(x[0],y[0],25,marker='p',color='r')
plt.text(x[-1], y[-1], ' S27', horizontalalignment='left', verticalalignment='center', fontsize=10, color='lightcoral', fontweight='bold')
x, y = m(stationLon[index_SESPB],stationLat[index_SESPB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], 'SESPB ', horizontalalignment='right', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_SWSPB],stationLat[index_SWSPB])
m.scatter(x,y,3,marker='o',color='r')
plt.text(x[-1], y[-1], 'SWSPB ', horizontalalignment='right', verticalalignment='center', fontsize=10, color='r', fontweight='bold')
x, y = m(stationLon[index_SS],stationLat[index_SS])
m.scatter(x,y,3,marker='o',color='lightcoral')
plt.text(x[0], y[0], 'SS ', horizontalalignment='right', verticalalignment='center', fontsize=10, color='lightcoral', fontweight='bold')


#m.plot(988291, 3.89471e6, '.r')
#lonpt, latpt = m(988291, 3.89471e6,inverse=True)
#plt.text(xpt+100000,ypt+100000,'(%5.1fW,%3.1fN)' % (lonpt,latpt))


for idx, rec in enumerate(records):
    if rec[-1] == '':
        continue
    else:
        coords = np.array(shapes[idx].points)
        x,y = m(coords[:,0], coords[:,1])
        m.plot(x,y,color='black')
        bbox = np.array(shapes[idx].bbox)
        x,y = m(np.mean([bbox[0], bbox[2]]), np.mean([bbox[1], bbox[3]]))                                                    
        plt.text(x, y, rec[-1], fontsize=10, color='black', fontweight='bold')




#### ---- Save Figure ---- ####
fig.set_size_inches(w=8, h=9)
fig.set_dpi(200)
fig.savefig(fig_name)
#plt.show()

